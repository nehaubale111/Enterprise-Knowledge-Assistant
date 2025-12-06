# backend/app/services/text_extraction.py
from pathlib import Path
from typing import List, Tuple

import pandas as pd
from pptx import Presentation
from pypdf import PdfReader
from docx import Document as DocxDocument


def extract_text_from_pdf(path: Path) -> List[Tuple[str, int]]:
    reader = PdfReader(str(path))
    pages = []
    for i, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        pages.append((text, i + 1))
    return pages


def extract_text_from_docx(path: Path) -> List[Tuple[str, int]]:
    doc = DocxDocument(str(path))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    return [(full_text, 1)]


def extract_text_from_pptx(path: Path) -> List[Tuple[str, int]]:
    prs = Presentation(str(path))
    texts = []
    for i, slide in enumerate(prs.slides):
        slide_text = []
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                slide_text.append(shape.text)
        texts.append(("\n".join(slide_text), i + 1))
    return texts


def extract_text_from_excel(path: Path) -> List[Tuple[str, int]]:
    df = pd.read_excel(path)
    text = df.to_string()
    return [(text, 1)]


def extract_text(path: Path, content_type: str) -> List[Tuple[str, int]]:
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        return extract_text_from_pdf(path)
    if suffix in (".docx",):
        return extract_text_from_docx(path)
    if suffix in (".pptx",):
        return extract_text_from_pptx(path)
    if suffix in (".xlsx", ".xls"):
        return extract_text_from_excel(path)

    # fallback: just read as text
    return [(path.read_text(encoding="utf-8", errors="ignore"), 1)]
